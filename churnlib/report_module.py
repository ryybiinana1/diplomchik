# churnlib/report_module.py
from __future__ import annotations

import base64
import io
from dataclasses import dataclass
from html import escape
from pathlib import Path
from typing import Any, Dict, List, Optional

import matplotlib.pyplot as plt
import numpy as np
import pandas as pd

from .economy_module import Scenario, expected_value, profit_curve, best_k


@dataclass(frozen=True)
class ReportConfig:
    out_dir: str
    top_k: int = 500
    value_col: str = "value_proxy"


def write_reports(df_test: pd.DataFrame, p: np.ndarray, scenarios: List[Scenario], cfg: ReportConfig) -> Dict[str, Any]:
    out = Path(cfg.out_dir)
    (out / "plots").mkdir(parents=True, exist_ok=True)
    (out / "tables").mkdir(parents=True, exist_ok=True)

    df = df_test.copy()
    df["p"] = np.asarray(p, dtype=float)
    V = df[cfg.value_col].astype(float).values

    base = next((s for s in scenarios if s.name == "base"), scenarios[0])
    df["EV_base"] = expected_value(df["p"].values, V, base)

    prio = df.sort_values("EV_base", ascending=False).head(cfg.top_k)
    prio_path = out / "tables" / "priority_list_topk.csv"
    prio.to_csv(prio_path, index=False)

    plt.figure(figsize=(8, 5))
    summary_rows = []
    base_best_k = None
    base_max_profit = None

    for sc in scenarios:
        curve = profit_curve(df["p"].values, V, sc)
        k, mx = best_k(curve)
        summary_rows.append({"scenario": sc.name, "best_k": k, "max_profit": mx})
        plt.plot(curve, label=f"{sc.name} (best_k={k})")
        if sc.name == "base":
            base_best_k = k
            base_max_profit = mx

    plt.title("Profit curve (cumulative EV for top-k)")
    plt.xlabel("k")
    plt.ylabel("Expected profit")
    plt.legend()
    plot_path = out / "plots" / "profit_curves.png"
    plt.tight_layout()
    plt.savefig(plot_path, dpi=150)
    plt.close()

    summary_path = out / "tables" / "profit_summary.csv"
    pd.DataFrame(summary_rows).to_csv(summary_path, index=False)

    return {
        "priority_csv": str(prio_path),
        "profit_plot": str(plot_path),
        "profit_summary": str(summary_path),
        "base_best_k": int(base_best_k) if base_best_k is not None else None,
        "base_max_profit": float(base_max_profit) if base_max_profit is not None else None,
    }


def _add_picture_if_exists(doc, path: Optional[str], width_inch: float = 6.0) -> None:
    if not path:
        return
    p = Path(path)
    if not p.exists():
        return
    try:
        from docx.shared import Inches
        doc.add_picture(str(p), width=Inches(width_inch))
    except Exception:
        return


def write_docx_report(
    out_path: str,
    template: str,
    params_used: Dict[str, Any],
    suitability: Dict[str, Any],
    metrics: Dict[str, Any],
    artifact_paths: Dict[str, Optional[str]],
    extra_feature_audit: Optional[List[Dict[str, Any]]] = None,
    mode: str = "single",
    experiment_rows: Optional[List[Dict[str, Any]]] = None,
    selection_metric: str = "pr_auc",
) -> Dict[str, Any]:
    """
    Полноценный .docx отчёт. Требует установленный python-docx. citeturn25search3
    """
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches, Pt, RGBColor
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    out_path = str(out_path)
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.55)
    section.left_margin = Inches(0.65)
    section.right_margin = Inches(0.65)

    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(9.5)

    def add_title(text: str) -> None:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(text)
        run.bold = True
        run.font.size = Pt(18)
        run.font.color.rgb = RGBColor(15, 23, 42)

    def add_subtitle(text: str) -> None:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(text)
        run.font.size = Pt(11)
        run.font.color.rgb = RGBColor(71, 85, 105)

    def add_table_from_rows(rows: List[Dict[str, Any]], cols: List[str]) -> None:
        if not rows or not cols:
            doc.add_paragraph("Данные недоступны.")
            return
        table = doc.add_table(rows=1, cols=len(cols))
        table.style = "Table Grid"
        for i, col in enumerate(cols):
            cell = table.rows[0].cells[i]
            cell.text = str(col)
            _set_cell_margins(cell)
            _shade_cell(cell, "DDEBFF")
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_before = Pt(2)
                paragraph.paragraph_format.space_after = Pt(2)
                for run in paragraph.runs:
                    run.bold = True
                    run.font.size = Pt(9.5)
                    run.font.color.rgb = RGBColor(30, 64, 175)
        for row in rows:
            cells = table.add_row().cells
            for i, col in enumerate(cols):
                value = row.get(col, "")
                cell = cells[i]
                _set_cell_margins(cell)
                cell.text = _fmt_num(value) if isinstance(value, (int, float, np.floating)) else str(value)
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.space_before = Pt(2)
                    paragraph.paragraph_format.space_after = Pt(2)
                    for run in paragraph.runs:
                        run.font.size = Pt(9)

    def _shade_cell(cell, fill: str) -> None:
        tc_pr = cell._tc.get_or_add_tcPr()
        shading = OxmlElement("w:shd")
        shading.set(qn("w:fill"), fill)
        tc_pr.append(shading)

    def _set_cell_margins(cell, top: int = 90, start: int = 110, bottom: int = 90, end: int = 110) -> None:
        tc_pr = cell._tc.get_or_add_tcPr()
        tc_mar = tc_pr.first_child_found_in("w:tcMar")
        if tc_mar is None:
            tc_mar = OxmlElement("w:tcMar")
            tc_pr.append(tc_mar)
        for margin_name, value in {"top": top, "start": start, "bottom": bottom, "end": end}.items():
            node = tc_mar.find(qn(f"w:{margin_name}"))
            if node is None:
                node = OxmlElement(f"w:{margin_name}")
                tc_mar.append(node)
            node.set(qn("w:w"), str(value))
            node.set(qn("w:type"), "dxa")

    def style_heading(text: str, level: int = 1):
        paragraph = doc.add_heading(text, level=level)
        for run in paragraph.runs:
            run.font.color.rgb = RGBColor(30, 64, 175)
        return paragraph

    exp_df = _normalize_experiment_rows(experiment_rows)
    best_models = _best_per_model(exp_df, selection_metric=selection_metric)
    best_row = best_models.head(1).iloc[0].to_dict() if not best_models.empty else {}
    best_model_name = str(best_row.get("model_kind") or params_used.get("model_kind") or "—")

    add_subtitle("Churn / Retention Studio · отчёт по моделям")
    add_title("ОТЧЁТ ПОСЛЕ ОБУЧЕНИЯ МОДЕЛИ")
    add_subtitle("Обучение и выбор модели прогнозирования оттока клиентов")

    hero = doc.add_table(rows=2, cols=4)
    hero.style = "Table Grid"
    hero.rows[0].cells[0].text = "Лучшая модель"
    hero.rows[0].cells[1].text = "ROC-AUC"
    hero.rows[0].cells[2].text = "PR-AUC"
    hero.rows[0].cells[3].text = "Brier score"
    hero.rows[1].cells[0].text = best_model_name
    hero.rows[1].cells[1].text = _fmt_num(best_row.get("roc_auc") or metrics.get("test_metrics_cal", {}).get("roc_auc"), 3)
    hero.rows[1].cells[2].text = _fmt_num(best_row.get("pr_auc") or metrics.get("test_metrics_cal", {}).get("pr_auc"), 3)
    hero.rows[1].cells[3].text = _fmt_num(best_row.get("brier") or metrics.get("test_metrics_cal", {}).get("brier"), 3)

    style_heading("1. Краткий итог эксперимента", level=1)
    doc.add_paragraph("Итог")
    doc.add_paragraph(
        "Система обработала датасет, проверила структуру входных данных, сопоставила пользовательские колонки "
        "с внутренней схемой, построила клиентские признаки и сравнила несколько моделей машинного обучения. "
        f"Лучшей моделью по основной метрике {selection_metric} выбрана {best_model_name}."
    )

    style_heading("2. Входной датасет", level=1)
    source_cols = params_used.get("input_source_columns") or []
    source_dtypes = params_used.get("input_source_dtypes") or {}
    data_rows = [
        {"Показатель": "Тип данных", "Значение": template},
        {"Показатель": "Количество колонок", "Значение": len(source_cols) if source_cols else "—"},
        {"Показатель": "Статус пригодности", "Значение": suitability.get("verdict", "—")},
        {"Показатель": "Целевая доля оттока", "Значение": _fmt_num(metrics.get("target_rate"), 3)},
    ]
    add_table_from_rows(data_rows, ["Показатель", "Значение"])

    style_heading("3. Сопоставление выбранных колонок", level=1)
    doc.add_paragraph("Перед обучением исходные колонки были приведены к внутренней схеме сервиса.")
    mapping_used = params_used.get("mapping_used") or {}
    mapping_rows = [
        {
            "Роль в системе": role,
            "Колонка в CSV": col or "—",
            "Внутреннее имя": role,
            "Статус": "использована" if col else "не задана",
        }
        for role, col in mapping_used.items()
    ]
    add_table_from_rows(mapping_rows, ["Роль в системе", "Колонка в CSV", "Внутреннее имя", "Статус"])

    style_heading("4. Использованные признаки", level=1)
    feature_cols = (params_used.get("training_schema") or {}).get("feature_cols") or params_used.get("feature_cols") or []
    extra_cols = params_used.get("extra_feature_cols") or []
    feature_rows = []
    for col in list(source_cols)[:14]:
        feature_rows.append({"Признак": col, "Тип данных": source_dtypes.get(col, "—"), "Комментарий": "исходная колонка"})
    for col in list(extra_cols)[:8]:
        feature_rows.append({"Признак": col, "Тип данных": source_dtypes.get(col, "—"), "Комментарий": "дополнительный признак"})
    for col in list(feature_cols)[:14]:
        feature_rows.append({"Признак": col, "Тип данных": "model feature", "Комментарий": "признак модели"})
    add_table_from_rows(feature_rows, ["Признак", "Тип данных", "Комментарий"])

    style_heading("5. Сравнение обученных моделей", level=1)
    test_raw = metrics.get("test_metrics_raw", {}) or {}
    test_cal = metrics.get("test_metrics_cal", {}) or {}
    wf = metrics.get("walk_forward_mean", {}) or {}
    if not exp_df.empty:
        doc.add_paragraph(
            "Модели сравнивались на отложенной тестовой выборке с учётом временного порядка данных. "
            f"Основной критерий выбора — {selection_metric}; дополнительно учитывались ROC-AUC, PR-AUC, Brier, Precision, Recall и F1."
        )
        ranked = best_models.copy()
        ranked.insert(0, "Место", range(1, len(ranked) + 1))
        ranked["Вердикт"] = ["Лучшая модель" if i == 0 else "Кандидат" for i in range(len(ranked))]
        add_table_from_rows(
            ranked.fillna("").to_dict(orient="records"),
            [c for c in ["Место", "model_kind", "roc_auc", "pr_auc", "brier", "precision", "recall", "f1", "Вердикт"] if c in ranked.columns],
        )

        cm_cols = {"model_kind", "cm_tn", "cm_fp", "cm_fn", "cm_tp"}
        if cm_cols.issubset(best_models.columns):
            doc.add_paragraph("Матрица ошибок для лучшей версии каждой модели:")
            cm_out = best_models.rename(columns={"model_kind": "Модель", "cm_tn": "TN", "cm_fp": "FP", "cm_fn": "FN", "cm_tp": "TP"})
            add_table_from_rows(cm_out.fillna("").to_dict(orient="records"), ["Модель", "TN", "FP", "FN", "TP"])

        time_best_df = _best_per_time_scenario(exp_df)
        if not time_best_df.empty:
            style_heading("5.1. Сравнение по временным сценариям", level=2)
            doc.add_paragraph(
                "Для каждой комбинации горизонта прогноза, окна истории и шага точки отсчёта выбран лучший алгоритм. "
                "Это позволяет учитывать каждый временной сценарий отдельно до итогового выбора модели."
            )
            add_table_from_rows(
                time_best_df.fillna("").to_dict(orient="records"),
                [c for c in ["horizon_days", "history_days", "step_days", "model_kind", "score", "roc_auc", "pr_auc", "brier"] if c in time_best_df.columns],
            )
    else:
        single_rows = [
            {"Метрика": "ROC-AUC", "Значение": test_cal.get("roc_auc")},
            {"Метрика": "PR-AUC", "Значение": test_cal.get("pr_auc")},
            {"Метрика": "Brier", "Значение": test_cal.get("brier")},
            {"Метрика": "Precision", "Значение": test_cal.get("precision")},
            {"Метрика": "Recall", "Значение": test_cal.get("recall")},
            {"Метрика": "F1", "Значение": test_cal.get("f1")},
            {"Метрика": "Walk-forward ROC-AUC", "Значение": wf.get("roc_auc")},
            {"Метрика": "Walk-forward PR-AUC", "Значение": wf.get("pr_auc")},
        ]
        add_table_from_rows(single_rows, ["Метрика", "Значение"])

    style_heading("6. Описание метрик", level=1)
    mf_rows = _metric_formula_rows()
    add_table_from_rows(mf_rows, ["Метрика", "Формула", "Описание"])

    style_heading("7. Графики качества моделей", level=1)
    chart_dir = Path(out_path).parent / "_docx_charts"
    if not best_models.empty:
        for idx, (metric_name, formula, desc, high) in enumerate(
            [
                ("roc_auc", "ROC-AUC = integral(TPR(FPR))", "Чем выше ROC-AUC, тем лучше модель разделяет клиентов с высоким и низким риском оттока.", True),
                ("pr_auc", "PR-AUC = integral(Precision(Recall))", "PR-AUC важна при дисбалансе классов и помогает оценить качество выявления клиентов из группы риска.", True),
                ("brier", "Brier = mean((p - y)^2)", "Brier score показывает точность вероятностного прогноза; меньшее значение лучше.", False),
            ],
            start=1,
        ):
            style_heading(f"7.{idx}. {metric_name.upper()}", level=2)
            doc.add_paragraph(f"Формула: {formula}")
            doc.add_paragraph(desc)
            chart_path = _save_model_metric_chart(best_models, metric_name, chart_dir / f"{metric_name}.png", higher_is_better=high)
            _add_picture_if_exists(doc, chart_path, width_inch=6.5)
            metric_table = best_models[["model_kind", metric_name]].rename(
                columns={"model_kind": "Модель", metric_name: metric_name.upper()}
            )
            add_table_from_rows(metric_table.fillna("").to_dict(orient="records"), ["Модель", metric_name.upper()])

    style_heading("7.4. Важность признаков", level=2)
    doc.add_paragraph(
        "График показывает, какие признаки сильнее всего влияли на прогноз лучшей модели. "
        "Это помогает объяснить результат пользователю и проверить бизнес-логику модели."
    )
    _add_picture_if_exists(doc, artifact_paths.get("shap_bar"), width_inch=6.5)
    _add_picture_if_exists(doc, artifact_paths.get("shap_beeswarm"), width_inch=6.5)

    style_heading("7.5. Экономический эффект", level=2)
    doc.add_paragraph(
        "Кривая ожидаемого эффекта показывает, сколько клиентов стоит включить в кампанию удержания. "
        "Ось X — top-k клиентов, ось Y — накопленный ожидаемый эффект. Линии соответствуют разным бизнес-сценариям."
    )
    _add_picture_if_exists(doc, artifact_paths.get("profit_plot"), width_inch=6.5)
    business_rows = [
        {
            "Показатель": "Рекомендуемый top-k",
            "Значение": metrics.get("business_metrics", {}).get("base_best_k"),
        },
        {
            "Показатель": "Максимальный ожидаемый эффект",
            "Значение": metrics.get("business_metrics", {}).get("base_max_profit"),
        },
    ]
    add_table_from_rows(business_rows, ["Показатель", "Значение"])

    if extra_feature_audit:
        style_heading("Дополнительные признаки", level=2)
        for a in extra_feature_audit:
            doc.add_paragraph(f"- {a.get('column')}: {a.get('verdict')} — {a.get('reason')}")

    style_heading("8. Обоснование выбора модели", level=1)
    doc.add_paragraph(f"Выбранная модель: {best_model_name}")
    doc.add_paragraph(
        f"{best_model_name} показала лучший результат по выбранной метрике {selection_metric} "
        "и сохраняет интерпретируемое качество по дополнительным метрикам. "
        "Модель рекомендована для скоринга клиентской базы и последующего расчёта экономического эффекта удержания."
    )

    style_heading("9. Ограничения и рекомендации", level=1)
    add_table_from_rows(
        [
            {"Тип": "Ограничение", "Описание": "Прогноз показывает вероятность оттока, а не гарантированный факт ухода клиента."},
            {"Тип": "Рекомендация", "Описание": "Переобучать модель на новых данных и контролировать изменение метрик качества."},
            {"Тип": "Рекомендация", "Описание": "Перед каждым скорингом проверять совпадение входных колонок с обучающей схемой."},
        ],
        ["Тип", "Описание"],
    )
    doc.add_paragraph(
        f"Итог: модель {best_model_name} рекомендована для скоринга клиентской базы, "
        "сегментации клиентов по уровню риска и дальнейшего расчёта экономического эффекта удержания."
    )

    Path(out_path).parent.mkdir(parents=True, exist_ok=True)
    doc.save(out_path)
    return {"docx_path": out_path}


def _safe_float(value: Any) -> Optional[float]:
    try:
        if value is None:
            return None
        return float(value)
    except Exception:
        return None


def _fmt_num(value: Any, digits: int = 4) -> str:
    num = _safe_float(value)
    if num is None:
        return "—"
    if abs(num) >= 1000:
        return f"{num:,.2f}".replace(",", " ")
    return f"{num:.{digits}f}"


def _fmt_int(value: Any) -> str:
    num = _safe_float(value)
    if num is None:
        return "—"
    return f"{int(num):,}".replace(",", " ")


def _normalize_experiment_rows(experiment_rows: Optional[List[Dict[str, Any]]]) -> pd.DataFrame:
    if not experiment_rows:
        return pd.DataFrame()
    df = pd.DataFrame(experiment_rows).copy()
    numeric_cols = [
        "score",
        "roc_auc",
        "pr_auc",
        "brier",
        "precision",
        "recall",
        "f1",
        "base_best_k",
        "base_max_profit",
        "cm_tn",
        "cm_fp",
        "cm_fn",
        "cm_tp",
    ]
    for col in numeric_cols:
        if col in df.columns:
            df[col] = pd.to_numeric(df[col], errors="coerce")
    if "status" in df.columns:
        df = df[df["status"].astype(str) == "ok"].copy()
    return df


def _best_per_model(exp_df: pd.DataFrame, selection_metric: str) -> pd.DataFrame:
    if exp_df.empty or "model_kind" not in exp_df.columns:
        return pd.DataFrame()
    work = exp_df.copy()
    score_col = "score" if "score" in work.columns else selection_metric
    if score_col not in work.columns:
        return pd.DataFrame()
    work = work.sort_values(score_col, ascending=False, na_position="last")
    best = work.groupby("model_kind", dropna=False, as_index=False).head(1)
    best = best.sort_values(score_col, ascending=False, na_position="last")
    return best.reset_index(drop=True)


def _confusion_grids_data_uri(best_models_df: pd.DataFrame) -> str:
    needed = {"model_kind", "cm_tn", "cm_fp", "cm_fn", "cm_tp"}
    if best_models_df.empty or not needed.issubset(best_models_df.columns):
        return ""
    rows = [r for _, r in best_models_df.iterrows() if pd.notna(r["cm_tn"]) and pd.notna(r["cm_tp"])]
    if not rows:
        return ""
    n = min(len(rows), 6)
    cols = 3
    rws = int(np.ceil(n / cols))
    fig, axes = plt.subplots(rws, cols, figsize=(12, 3.8 * rws))
    if not isinstance(axes, np.ndarray):
        axes = np.array([axes])
    axes = axes.reshape(-1)
    for idx in range(len(axes)):
        ax = axes[idx]
        if idx >= n:
            ax.axis("off")
            continue
        row = rows[idx]
        mat = np.array([[row["cm_tn"], row["cm_fp"]], [row["cm_fn"], row["cm_tp"]]], dtype=float)
        im = ax.imshow(mat, cmap="Blues")
        ax.set_xticks([0, 1], ["Pred 0", "Pred 1"])
        ax.set_yticks([0, 1], ["True 0", "True 1"])
        ax.set_title(str(row.get("model_kind", "model")))
        for i in range(2):
            for j in range(2):
                ax.text(j, i, _fmt_int(mat[i, j]), ha="center", va="center", color="#111827", fontsize=9)
        fig.colorbar(im, ax=ax, fraction=0.046, pad=0.04)
    fig.suptitle("Confusion matrix для лучшей версии каждой модели", fontsize=12)
    uri = _plt_to_data_uri(fig)
    plt.close(fig)
    return uri


def _metric_formula_rows() -> List[Dict[str, str]]:
    return [
        {
            "Метрика": "Precision",
            "Формула": "Precision = TP / (TP + FP)",
            "Описание": "Доля верно предсказанных оттоков среди всех объектов, которые модель пометила как отток.",
        },
        {
            "Метрика": "Recall",
            "Формула": "Recall = TP / (TP + FN)",
            "Описание": "Доля найденных оттоков среди всех реальных оттоков.",
        },
        {
            "Метрика": "F1",
            "Формула": "F1 = 2 * Precision * Recall / (Precision + Recall)",
            "Описание": "Баланс между точностью и полнотой.",
        },
        {
            "Метрика": "ROC-AUC",
            "Формула": "AUC = integral(TPR(FPR))",
            "Описание": "Способность модели отделять классы по всем порогам.",
        },
        {
            "Метрика": "PR-AUC",
            "Формула": "AUC = integral(Precision(Recall))",
            "Описание": "Ключевая метрика для несбалансированных классов; показывает качество ранжирования риска оттока.",
        },
        {
            "Метрика": "Brier",
            "Формула": "Brier = mean((p - y)^2)",
            "Описание": "Среднеквадратичная ошибка вероятностей (меньше — лучше).",
        },
    ]


def _best_per_time_scenario(exp_df: pd.DataFrame) -> pd.DataFrame:
    need = {"horizon_days", "history_days", "step_days", "score"}
    if exp_df.empty or not need.issubset(exp_df.columns):
        return pd.DataFrame()
    work = exp_df.copy()
    work = work.sort_values("score", ascending=False, na_position="last")
    best = work.groupby(["horizon_days", "history_days", "step_days"], dropna=False, as_index=False).head(1)
    best = best.sort_values(["horizon_days", "history_days", "step_days", "score"], ascending=[True, True, True, False])
    return best.reset_index(drop=True)


def _save_model_metric_chart(df: pd.DataFrame, metric: str, out_path: Path, *, higher_is_better: bool = True) -> Optional[str]:
    if df.empty or metric not in df.columns or "model_kind" not in df.columns:
        return None
    chart_df = df[["model_kind", metric]].copy()
    chart_df[metric] = pd.to_numeric(chart_df[metric], errors="coerce")
    chart_df = chart_df.dropna(subset=[metric])
    if chart_df.empty:
        return None
    chart_df = chart_df.sort_values(metric, ascending=not higher_is_better)
    fig, ax = plt.subplots(figsize=(7.2, 3.6))
    ax.bar(chart_df["model_kind"].astype(str), chart_df[metric], color="#2563eb")
    ax.set_title(f"Сравнение моделей по {metric}")
    ax.set_ylabel(metric)
    ax.tick_params(axis="x", rotation=25)
    fig.tight_layout()
    out_path.parent.mkdir(parents=True, exist_ok=True)
    fig.savefig(out_path, dpi=160, bbox_inches="tight")
    plt.close(fig)
    return str(out_path)


def _save_scenario_max_effect_chart(df: pd.DataFrame, out_path: Path) -> Optional[str]:
    if df.empty or not {"scenario", "max_profit"}.issubset(df.columns):
        return None
    chart_df = df[["scenario", "max_profit"]].copy()
    chart_df["max_profit"] = pd.to_numeric(chart_df["max_profit"], errors="coerce")
    chart_df = chart_df.dropna(subset=["max_profit"])
    if chart_df.empty:
        return None
    fig, ax = plt.subplots(figsize=(7.2, 3.6))
    ax.bar(chart_df["scenario"].astype(str), chart_df["max_profit"], color="#2563eb")
    ax.set_title("Максимальный эффект по сценариям")
    ax.set_ylabel("max_profit")
    ax.set_xlabel("Сценарий")
    fig.tight_layout()
    out_path.parent.mkdir(parents=True, exist_ok=True)
    fig.savefig(out_path, dpi=160, bbox_inches="tight")
    plt.close(fig)
    return str(out_path)

def _img_data_uri(path: Optional[str]) -> str:
    if not path:
        return ""
    p = Path(path)
    if not p.exists() or not p.is_file():
        return ""
    ext = p.suffix.lower()
    mime = {
        ".png": "image/png",
        ".jpg": "image/jpeg",
        ".jpeg": "image/jpeg",
        ".gif": "image/gif",
        ".webp": "image/webp",
    }.get(ext)
    if not mime:
        return ""
    payload = base64.b64encode(p.read_bytes()).decode("ascii")
    return f"data:{mime};base64,{payload}"


def _plt_to_data_uri(fig: plt.Figure) -> str:
    buf = io.BytesIO()
    fig.savefig(buf, format="png", dpi=140, bbox_inches="tight")
    buf.seek(0)
    payload = base64.b64encode(buf.read()).decode("ascii")
    return f"data:image/png;base64,{payload}"


def _html_table(rows: List[Dict[str, Any]], columns: List[str]) -> str:
    if not rows:
        return "<p>Данные недоступны.</p>"
    head = "".join(f"<th>{escape(str(c))}</th>" for c in columns)
    body_rows = []
    for row in rows:
        tds = "".join(f"<td>{escape(str(row.get(c, '—')))}</td>" for c in columns)
        body_rows.append(f"<tr>{tds}</tr>")
    body = "".join(body_rows)
    return f"<table><thead><tr>{head}</tr></thead><tbody>{body}</tbody></table>"


def _page_html(title: str, blocks: List[str]) -> str:
    style = """
    <style>
      body { font-family: -apple-system, BlinkMacSystemFont, 'Segoe UI', Arial, sans-serif; margin: 24px auto; max-width: 1100px; color: #0f172a; line-height: 1.45; background: #f8fafc; }
      h1, h2, h3 { color: #0f172a; margin: 0 0 10px 0; }
      h1 { font-size: 28px; margin-bottom: 16px; }
      h2 { margin-top: 22px; font-size: 20px; }
      h3 { margin-top: 14px; font-size: 16px; }
      .muted { color: #6b7280; }
      .card { border: 1px solid #e2e8f0; border-radius: 14px; padding: 16px; margin: 12px 0; background: #fff; box-shadow: 0 1px 2px rgba(2,6,23,.04); }
      .kpis { display: flex; flex-wrap: wrap; gap: 10px; }
      .kpi { border: 1px solid #e2e8f0; border-radius: 10px; padding: 10px 12px; min-width: 180px; background: #f8fafc; }
      .kpi .v { font-size: 20px; font-weight: 700; margin-top: 6px; }
      table { border-collapse: collapse; width: 100%; margin: 10px 0; }
      th, td { border: 1px solid #e2e8f0; padding: 8px; text-align: left; vertical-align: top; }
      th { background: #f1f5f9; }
      img { max-width: 100%; border: 1px solid #e2e8f0; border-radius: 10px; margin: 8px 0; background: #fff; }
      ul { margin: 8px 0 8px 20px; }
      code { background: #f1f5f9; padding: 1px 4px; border-radius: 4px; }
    </style>
    """
    return (
        "<!doctype html><html><head><meta charset='utf-8'/>"
        f"<title>{escape(title)}</title>{style}</head><body>"
        f"<h1>{escape(title)}</h1>"
        + "".join(blocks)
        + "</body></html>"
    )


def write_training_html_report(
    out_path: str,
    template: str,
    params_used: Dict[str, Any],
    suitability: Dict[str, Any],
    metrics: Dict[str, Any],
    artifact_paths: Dict[str, Optional[str]],
    mode: str = "single",
    experiment_rows: Optional[List[Dict[str, Any]]] = None,
    selection_metric: str = "pr_auc",
) -> Dict[str, Any]:
    target_rate = _fmt_num(metrics.get("target_rate"), digits=4)
    best_k = metrics.get("business_metrics", {}).get("base_best_k")
    max_profit = _fmt_num(metrics.get("business_metrics", {}).get("base_max_profit"), digits=2)
    test_cal = metrics.get("test_metrics_cal", {}) or {}
    wf = metrics.get("walk_forward_mean", {}) or {}
    reasons = suitability.get("reasons") or []
    exp_df = _normalize_experiment_rows(experiment_rows)
    best_models_df = _best_per_model(exp_df, selection_metric=selection_metric)

    metric_rows = [
        {"Метрика": "ROC-AUC (test, calibrated)", "Значение": _fmt_num(test_cal.get("roc_auc"))},
        {"Метрика": "PR-AUC (test, calibrated)", "Значение": _fmt_num(test_cal.get("pr_auc"))},
        {"Метрика": "Brier (test, calibrated)", "Значение": _fmt_num(test_cal.get("brier"))},
        {"Метрика": "Precision", "Значение": _fmt_num(test_cal.get("precision"))},
        {"Метрика": "Recall", "Значение": _fmt_num(test_cal.get("recall"))},
        {"Метрика": "F1", "Значение": _fmt_num(test_cal.get("f1"))},
        {"Метрика": "Walk-forward ROC-AUC (mean)", "Значение": _fmt_num(wf.get("roc_auc"))},
        {"Метрика": "Walk-forward PR-AUC (mean)", "Значение": _fmt_num(wf.get("pr_auc"))},
        {"Метрика": "Walk-forward Brier (mean)", "Значение": _fmt_num(wf.get("brier"))},
    ]

    blocks = [
        f"<p class='muted'>Тип данных: <b>{escape(str(template))}</b> · режим: <b>{escape(str(mode))}</b> · метрика отбора: <b>{escape(str(selection_metric))}</b></p>",
        "<div class='card'><h2>Итог</h2>"
        "<div class='kpis'>"
        f"<div class='kpi'><div>Target rate</div><div class='v'>{target_rate}</div></div>"
        f"<div class='kpi'><div>Лучший top-k (base)</div><div class='v'>{escape(str(best_k))}</div></div>"
        f"<div class='kpi'><div>Макс. ожидаемый эффект</div><div class='v'>{max_profit}</div></div>"
        f"<div class='kpi'><div>PR-AUC</div><div class='v'>{_fmt_num(test_cal.get('pr_auc'))}</div></div>"
        "</div></div>",
        "<div class='card'><h2>Параметры обучения</h2><ul>"
        + "".join(
            f"<li><b>{escape(str(k))}</b>: {escape(str(v))}</li>"
            for k, v in params_used.items()
            if k in {"model_name", "model_kind", "horizon_days", "history_days", "step_days", "calibration"}
        )
        + "</ul></div>",
        "<div class='card'><h2>Метрики финальной модели</h2>"
        + _html_table(metric_rows, ["Метрика", "Значение"])
        + "</div>",
    ]

    if not best_models_df.empty:
        cmp_cols = [c for c in ["model_kind", "experiment", "score", "roc_auc", "pr_auc", "brier", "precision", "recall", "f1", "base_max_profit"] if c in best_models_df.columns]
        blocks.append(
            "<div class='card'><h2>Сравнение всех моделей</h2>"
            "<p>Для каждой модели выбрана лучшая конфигурация. Основа выбора — selection metric, но в таблице показаны все ключевые метрики.</p>"
            + _html_table(best_models_df.fillna("").to_dict(orient="records"), cmp_cols)
            + "</div>"
        )

        if {"model_kind", "score"}.issubset(best_models_df.columns):
            bdf = best_models_df.copy()
            bdf = bdf.sort_values("score", ascending=False, na_position="last")
            fig, ax = plt.subplots(figsize=(9, 4.6))
            ax.barh(bdf["model_kind"].astype(str)[::-1], bdf["score"][::-1], color="#2563eb")
            ax.set_title(f"Ранжирование моделей по {selection_metric}")
            ax.set_xlabel("score")
            rank_uri = _plt_to_data_uri(fig)
            plt.close(fig)
            blocks.append("<div class='card'><h2>Ранжирование моделей</h2>" + f"<img src='{rank_uri}' alt='Model ranking'/>" + "</div>")

        cm_uri = _confusion_grids_data_uri(best_models_df)
        if cm_uri:
            blocks.append(
                "<div class='card'><h2>Confusion matrix для моделей</h2>"
                "<p>Показаны матрицы ошибок для лучшей версии каждой модели при рабочем пороге. "
                "Диагональ (TN+TP) должна быть максимально большой, недиагональные клетки (FP+FN) — минимальными.</p>"
                f"<img src='{cm_uri}' alt='Confusion matrices'/>"
                "</div>"
            )

    time_best_df = _best_per_time_scenario(exp_df)
    if not time_best_df.empty:
        cols = [
            c
            for c in [
                "horizon_days",
                "history_days",
                "step_days",
                "model_kind",
                "score",
                "roc_auc",
                "pr_auc",
                "brier",
                "precision",
                "recall",
                "f1",
                "base_max_profit",
            ]
            if c in time_best_df.columns
        ]
        blocks.append(
            "<div class='card'><h2>Сравнение по временным сценариям</h2>"
            "<p>Каждая комбинация горизонта/окна/шага анализируется отдельно, "
            "и для неё выбирается лучший вариант по метрике отбора.</p>"
            + _html_table(time_best_df.fillna("").to_dict(orient="records"), cols)
            + "</div>"
        )

    formula_rows = _metric_formula_rows()
    blocks.append(
        "<div class='card'><h2>Формулы и трактовка метрик</h2>"
        "<p>Перед анализом графиков ниже приведены формулы и краткие пояснения по ключевым метрикам.</p>"
        + _html_table(formula_rows, ["Метрика", "Формула", "Описание"])
        + "</div>"
    )

    verdict = suitability.get("verdict", "unknown")
    blocks.append(
        "<div class='card'><h2>Проверка данных</h2>"
        f"<p>Вердикт пригодности: <b>{escape(str(verdict))}</b>.</p>"
        + ("<ul>" + "".join(f"<li>{escape(str(r))}</li>" for r in reasons) + "</ul>" if reasons else "")
        + "</div>"
    )

    profit_uri = _img_data_uri(artifact_paths.get("profit_plot"))
    cal_uri = _img_data_uri(artifact_paths.get("calibration_plot"))
    shap_uri = _img_data_uri(artifact_paths.get("shap_beeswarm"))
    if profit_uri or cal_uri or shap_uri:
        visuals = "<div class='card'><h2>Графики и интерпретация</h2>"
        if profit_uri:
            visuals += (
                "<h3>Кривая ожидаемого эффекта (profit curve)</h3>"
                "<p>График построен для итоговой выбранной модели. По оси X — размер кампании top-k, "
                "по оси Y — кумулятивный ожидаемый эффект. Линии соответствуют сценариям "
                "<code>conservative</code>, <code>base</code>, <code>optimistic</code> и меняются при изменении "
                "параметров margin/cost/success и качества вероятностей модели.</p>"
                f"<img src='{profit_uri}' alt='Profit curve'/>"
            )
        if cal_uri:
            visuals += (
                "<h3>Калибровка вероятностей</h3>"
                "<p>Показывает, насколько прогнозируемые вероятности риска соответствуют фактической частоте оттока.</p>"
                f"<img src='{cal_uri}' alt='Calibration curve'/>"
            )
        if shap_uri:
            visuals += (
                "<h3>Глобальные причины (SHAP)</h3>"
                "<p>SHAP показывает вклад признаков в общий риск. "
                "Чем дальше признак от центра по оси SHAP, тем сильнее его влияние.</p>"
                f"<img src='{shap_uri}' alt='SHAP beeswarm'/>"
            )
        visuals += "</div>"
        blocks.append(visuals)

    if not exp_df.empty and {"weights_model", "weights_calibrator"}.issubset(exp_df.columns):
        weight_rows = exp_df[["experiment", "model_kind", "weights_model", "weights_calibrator"]].fillna("").to_dict(orient="records")
        blocks.append(
            "<div class='card'><h2>Сохранённые веса моделей</h2>"
            "<p>Ниже веса всех обученных вариантов. Эти файлы включены в ZIP.</p>"
            + _html_table(weight_rows, ["experiment", "model_kind", "weights_model", "weights_calibrator"])
            + "</div>"
        )

    blocks.append(
        "<div class='card'><h2>Рекомендации</h2><ul>"
        "<li>Запускать удержание с top-k из базового сценария, затем масштабировать по фактическому uplift.</li>"
        "<li>Регулярно обновлять модель на новых данных и пересчитывать экономические допущения.</li>"
        "<li>Использовать сравнение моделей не только по PR-AUC, но и по Precision/Recall/F1 и confusion matrix.</li>"
        "</ul></div>"
    )

    html = _page_html("Единый отчёт по обучению модели", blocks)
    out_p = Path(out_path)
    out_p.parent.mkdir(parents=True, exist_ok=True)
    out_p.write_text(html, encoding="utf-8")
    return {"html_path": str(out_p)}


def write_scoring_html_report(
    out_path: str,
    model_info: Dict[str, Any],
    n_scored: int,
    business_summary: Dict[str, Any],
    scenario_rows: List[Dict[str, Any]],
    top_clients_preview: List[Dict[str, Any]],
    score_mapping: Dict[str, str],
) -> Dict[str, Any]:
    scenario_df = pd.DataFrame(scenario_rows or [])
    top_df = pd.DataFrame(top_clients_preview or [])

    blocks: List[str] = [
        (
            "<p class='muted'>"
            f"Модель: <b>{escape(str(model_info.get('model_name', '—')))}</b> · "
            f"Шаблон: <b>{escape(str(model_info.get('template', '—')))}</b> · "
            f"Алгоритм: <b>{escape(str(model_info.get('model_kind', '—')))}</b> · "
            f"Горизонт: <b>{escape(str(model_info.get('horizon_days', '—')))} дней</b>"
            "</p>"
        ),
        "<div class='card'><h2>Итог прогноза</h2>"
        "<div class='kpis'>"
        f"<div class='kpi'><div>Оценено строк</div><div class='v'>{escape(str(n_scored))}</div></div>"
        f"<div class='kpi'><div>Клиентов с EV &gt; 0</div><div class='v'>{escape(str(business_summary.get('clients_with_positive_ev', '—')))}</div></div>"
        f"<div class='kpi'><div>Суммарный положительный EV</div><div class='v'>{escape(_fmt_num(business_summary.get('total_positive_ev'), digits=2))}</div></div>"
        f"<div class='kpi'><div>Рекомендуемый top-k</div><div class='v'>{escape(str(business_summary.get('best_k', '—')))}</div></div>"
        "</div></div>",
    ]

    if not scenario_df.empty:
        cols = [c for c in ["scenario", "margin", "cost", "success", "clients_with_positive_ev", "best_k", "max_profit"] if c in scenario_df.columns]
        blocks.append(
            "<div class='card'><h2>Сценарный анализ</h2>"
            "<p>Таблица показывает устойчивость решения при разных бизнес-допущениях.</p>"
            + _html_table(scenario_df.fillna("").to_dict(orient="records"), cols)
            + "</div>"
        )
        if {"scenario", "max_profit"}.issubset(scenario_df.columns):
            chart_df = scenario_df.copy()
            chart_df["max_profit"] = pd.to_numeric(chart_df["max_profit"], errors="coerce")
            chart_df = chart_df.dropna(subset=["max_profit"])
            if not chart_df.empty:
                fig, ax = plt.subplots(figsize=(7.5, 4.2))
                ax.bar(chart_df["scenario"].astype(str), chart_df["max_profit"], color="#0f766e")
                ax.set_title("Максимальный ожидаемый эффект по сценариям")
                ax.set_ylabel("max_profit")
                chart_uri = _plt_to_data_uri(fig)
                plt.close(fig)
                blocks.append("<div class='card'><h2>График сценариев</h2>" + f"<img src='{chart_uri}' alt='Scenario max profit'/>" + "</div>")

    if not top_df.empty:
        col_order = [c for c in ["entity_id", "customer_id", "client_id", "account_id", "risk_segment", "priority", "EV_base", "p_calibrated", "reason_1"] if c in top_df.columns]
        if not col_order:
            col_order = list(top_df.columns[:8])
        blocks.append(
            "<div class='card'><h2>Top-клиенты для базового сценария</h2>"
            "<p>Ниже первые строки приоритетного списка, ранжированного по EV_base.</p>"
            + _html_table(top_df.fillna("").to_dict(orient="records"), col_order)
            + "</div>"
        )

    if score_mapping:
        map_rows = [{"Колонка в обучении": src, "Колонка в файле прогноза": dst} for src, dst in score_mapping.items()]
        blocks.append(
            "<div class='card'><h2>Сопоставление колонок</h2>"
            "<p>Маппинг, использованный для прогноза этой модели.</p>"
            + _html_table(map_rows, ["Колонка в обучении", "Колонка в файле прогноза"])
            + "</div>"
        )

    blocks.append(
        "<div class='card'><h2>Пояснение к соседним CSV-файлам</h2><ul>"
        "<li><code>scored_clients.csv</code> — полный скоринг по всем объектам, включая вероятности и EV.</li>"
        "<li><code>retention_priority_list.csv</code> — базовый ранжированный список для кампании удержания.</li>"
        "<li><code>scenario_summary.csv</code> — сводка по сценариям с best_k и max_profit.</li>"
        "<li><code>retention_priority_list_conservative.csv</code> — приоритеты для осторожного сценария.</li>"
        "<li><code>retention_priority_list_base.csv</code> — приоритеты для базового сценария.</li>"
        "<li><code>retention_priority_list_optimistic.csv</code> — приоритеты для оптимистичного сценария.</li>"
        "</ul></div>"
    )

    blocks.append(
        "<div class='card'><h2>Управленческий вывод</h2>"
        f"<p>{escape(str(business_summary.get('scenario_comment') or 'Начинать рекомендуется с клиентов из top-k базового сценария, затем масштабировать кампанию по результатам пилота.'))}</p>"
        "</div>"
    )

    html = _page_html("Единый отчёт по прогнозу и экономическому эффекту", blocks)
    out_p = Path(out_path)
    out_p.parent.mkdir(parents=True, exist_ok=True)
    out_p.write_text(html, encoding="utf-8")
    return {"html_path": str(out_p)}


def _describe_csv_column(column: str) -> str:
    descriptions = {
        "entity_id": "Идентификатор клиента, аккаунта или объекта, для которого рассчитан прогноз. Используется как ключ строки результата.",
        "customer_id": "Идентификатор клиента из исходного CSV, если он был сохранён в результате.",
        "client_id": "Идентификатор клиента из исходного CSV, если в данных использовалось такое название.",
        "account_id": "Идентификатор аккаунта из исходного CSV для подписочной модели.",
        "user_id": "Идентификатор пользователя из исходного CSV для событийной модели.",
        "subject_id": "Идентификатор субъекта/пользователя из событийных данных.",
        "anchor_time": "Дата точки отсчёта, на которую собрана история клиента и рассчитан прогноз.",
        "target": "Фактический обучающий target, если он сохранён в диагностической выгрузке: 1 — отток/нет будущей активности, 0 — активность сохранилась.",
        "p_raw": "Вероятность оттока, которую выдала модель до калибровки вероятностей.",
        "p_cal": "Откалиброванная вероятность оттока на тестовой/диагностической выборке.",
        "p_calibrated": "Итоговая вероятность оттока после калибровки. Чем выше значение, тем выше риск.",
        "p": "Вероятность оттока, использованная при расчёте ожидаемого эффекта удержания.",
        "pred": "Бинарное решение по рабочему порогу: 1 — модель относит объект к риску оттока, 0 — не относит.",
        "risk_segment": "Группа риска, присвоенная по p_calibrated: low, medium, high или critical.",
        "value_proxy": "Оценка ценности клиента для экономического расчёта. Для транзакций обычно основана на сумме покупок, для подписок — на MRR, для событий — на активности.",
        "EV": "Ожидаемый эффект удержания в базовом сценарии. Дублирует EV_base для удобства сортировки.",
        "EV_base": "Ожидаемый эффект удержания в базовом сценарии: вероятность оттока × ценность × success × margin − cost.",
        "EV_conservative": "Ожидаемый эффект в осторожном сценарии с более жёсткими допущениями по стоимости и успеху удержания.",
        "EV_optimistic": "Ожидаемый эффект в оптимистичном сценарии с более благоприятными бизнес-допущениями.",
        "priority": "Приоритет для кампании удержания: высокий, средний или низкий в зависимости от EV и позиции в рейтинге.",
        "recommended_action": "Текстовая рекомендация: включить в кампанию, рассмотреть при наличии ресурса или не приоритизировать.",
        "priority_scenario": "Экономический сценарий, для которого сформирован данный приоритетный список.",
        "scenario_ev": "Ожидаемый эффект, по которому отсортирован сценарный список клиентов.",
        "reason_1": "Главная автоматически сформированная причина риска, например давняя активность или редкие покупки.",
        "reason_2": "Вторая причина риска, если для клиента найдено несколько объяснений.",
        "reason_3": "Третья причина риска, если она доступна.",
        "scenario": "Название сценария экономики удержания: conservative, base или optimistic.",
        "margin": "Доля ценности клиента, которую бизнес условно сохраняет при успешном удержании.",
        "cost": "Стоимость одного удерживающего контакта или действия на клиента.",
        "success": "Ожидаемая вероятность успешного удержания клиента после контакта.",
        "clients_with_positive_ev": "Сколько клиентов в сценарии имеют положительный ожидаемый эффект удержания.",
        "best_k": "Оптимальное количество клиентов для кампании в этом сценарии: top-k с максимальным суммарным EV.",
        "max_profit": "Максимальный суммарный ожидаемый эффект, достигаемый при best_k.",
        "total_positive_ev": "Сумма всех положительных EV в сценарии без ограничения на best_k.",
        "mean_ev_top20": "Средний ожидаемый эффект среди первых 20 клиентов рейтинга.",
        "frequency_tx": "Количество покупок/транзакций клиента в историческом окне перед прогнозом.",
        "monetary": "Суммарная сумма покупок клиента в историческом окне.",
        "amount_mean": "Средняя сумма одной покупки клиента в историческом окне.",
        "amount_std": "Стандартное отклонение суммы покупок клиента; показывает разброс чеков.",
        "recency_days": "Сколько дней прошло от последней активности/покупки до точки прогноза.",
        "customer_lifetime_days": "Сколько дней клиент наблюдается в данных до точки прогноза: от первой активности до anchor_time.",
        "interpurchase_mean_days": "Средний интервал в днях между покупками клиента.",
        "interpurchase_std_days": "Разброс интервалов между покупками; показывает регулярность или нерегулярность активности.",
        "qty_sum": "Суммарное количество купленных единиц товара в историческом окне.",
        "qty_mean": "Среднее количество единиц товара в одной покупке.",
        "unit_price_mean": "Средняя цена единицы товара в покупках клиента.",
        "unit_price_std": "Разброс цены единицы товара в покупках клиента.",
        "item_nunique": "Количество уникальных товаров, купленных клиентом.",
        "country_nunique": "Количество уникальных стран/географий, встречавшихся в истории клиента.",
        "mrr_sum": "Суммарный MRR аккаунта в историческом окне для подписочных данных.",
        "mrr_last": "Последний известный MRR аккаунта перед точкой прогноза.",
        "event_count": "Количество событий пользователя в историческом окне для событийных данных.",
    }
    if column in descriptions:
        return descriptions[column]
    if column.startswith("extra__"):
        parts = column.split("__")
        source = parts[1] if len(parts) > 1 else "дополнительной колонки"
        suffix = parts[2] if len(parts) > 2 else ""
        suffix_descriptions = {
            "mean": "среднее значение",
            "std": "разброс значений",
            "min": "минимальное значение",
            "max": "максимальное значение",
            "sum": "сумма значений",
            "last": "последнее известное значение",
            "nunique": "количество уникальных значений",
            "entropy": "разнообразие категорий: чем выше, тем менее однотипны значения",
            "top1_share": "доля самого частого значения",
            "missing_share": "доля пропусков",
            "code_mean": "средний числовой код категории",
            "code_max": "максимальный числовой код категории",
            "code_last": "последний числовой код категории",
            "age_mean_days": "средний возраст даты в днях относительно точки прогноза",
            "age_std_days": "разброс возраста даты в днях",
            "age_min_days": "минимальный возраст даты в днях",
            "age_max_days": "максимальный возраст даты в днях",
        }
        if suffix.startswith("share_"):
            value = suffix.replace("share_", "", 1)
            return f"Доля строк клиента, где дополнительная колонка `{source}` имела значение `{value}`."
        if suffix.startswith("last_is_"):
            value = suffix.replace("last_is_", "", 1)
            return f"Флаг того, что последнее значение дополнительной колонки `{source}` равно `{value}`."
        detail = suffix_descriptions.get(suffix, f"агрегат `{suffix}`")
        return f"Агрегат по дополнительной колонке `{source}`: {detail} в истории клиента до прогноза."
    if column.startswith("EV_"):
        scenario = column.replace("EV_", "", 1)
        return f"Ожидаемый эффект удержания для сценария `{scenario}`."
    if column.endswith("_raw"):
        base = column.removesuffix("_raw")
        return f"Исходное значение `{base}` до финального преобразования или калибровки."
    if column.endswith("_cal"):
        base = column.removesuffix("_cal")
        return f"Откалиброванное значение `{base}` после постобработки модели."
    return (
        f"Колонка `{column}` сохранена из подготовленного набора признаков модели. "
        "Она содержит числовую характеристику поведения клиента, рассчитанную на историческом окне перед прогнозом."
    )


def _csv_column_dictionary(columns: List[str]) -> List[Dict[str, str]]:
    return [{"Колонка": col, "Описание": _describe_csv_column(col)} for col in columns]


def write_scoring_docx_report(
    out_path: str,
    model_info: Dict[str, Any],
    n_scored: int,
    business_summary: Dict[str, Any],
    scenario_rows: List[Dict[str, Any]],
    top_clients_preview: List[Dict[str, Any]],
    score_mapping: Dict[str, str],
    csv_paths: Dict[str, str],
) -> Dict[str, Any]:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor

    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.6)
    section.bottom_margin = Inches(0.6)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)
    doc.styles["Normal"].font.name = "Arial"
    doc.styles["Normal"].font.size = Pt(9.5)

    def shade_cell(cell, fill: str = "DDEBFF") -> None:
        tc_pr = cell._tc.get_or_add_tcPr()
        shading = OxmlElement("w:shd")
        shading.set(qn("w:fill"), fill)
        tc_pr.append(shading)

    def set_cell_margins(cell) -> None:
        tc_pr = cell._tc.get_or_add_tcPr()
        tc_mar = tc_pr.first_child_found_in("w:tcMar")
        if tc_mar is None:
            tc_mar = OxmlElement("w:tcMar")
            tc_pr.append(tc_mar)
        for name, value in {"top": 90, "start": 110, "bottom": 90, "end": 110}.items():
            node = tc_mar.find(qn(f"w:{name}"))
            if node is None:
                node = OxmlElement(f"w:{name}")
                tc_mar.append(node)
            node.set(qn("w:w"), str(value))
            node.set(qn("w:type"), "dxa")

    def add_heading(text: str, level: int = 1) -> None:
        paragraph = doc.add_heading(text, level=level)
        for run in paragraph.runs:
            run.font.color.rgb = RGBColor(30, 64, 175)

    def add_table(rows: List[Dict[str, Any]], cols: List[str]) -> None:
        if not rows:
            doc.add_paragraph("Данные недоступны.")
            return
        table = doc.add_table(rows=1, cols=len(cols))
        table.style = "Table Grid"
        for i, col in enumerate(cols):
            cell = table.rows[0].cells[i]
            cell.text = col
            shade_cell(cell)
            set_cell_margins(cell)
            for p in cell.paragraphs:
                for run in p.runs:
                    run.bold = True
                    run.font.color.rgb = RGBColor(30, 64, 175)
        for row in rows:
            cells = table.add_row().cells
            for i, col in enumerate(cols):
                cells[i].text = str(row.get(col, "—"))
                set_cell_margins(cells[i])

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("ОТЧЁТ ПО РЕЗУЛЬТАТАМ ПРОГНОЗА")
    run.bold = True
    run.font.size = Pt(18)
    run.font.color.rgb = RGBColor(15, 23, 42)
    subtitle = doc.add_paragraph("Churn / Retention Studio · прогноз риска и экономического эффекта")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER

    add_heading("1. Краткий итог", level=1)
    add_table(
        [
            {"Показатель": "Модель", "Значение": model_info.get("model_name", "—")},
            {"Показатель": "Алгоритм", "Значение": model_info.get("model_kind", "—")},
            {"Показатель": "Горизонт прогноза", "Значение": model_info.get("horizon_days", "—")},
            {"Показатель": "Оценено строк", "Значение": n_scored},
            {"Показатель": "Клиентов с EV > 0", "Значение": business_summary.get("clients_with_positive_ev", "—")},
            {"Показатель": "Рекомендуемый top-k", "Значение": business_summary.get("best_k", "—")},
            {"Показатель": "Суммарный положительный EV", "Значение": _fmt_num(business_summary.get("total_positive_ev"), 2)},
        ],
        ["Показатель", "Значение"],
    )

    add_heading("2. Сценарии экономического эффекта", level=1)
    doc.add_paragraph(
        "Сценарии показывают, насколько решение устойчиво к разным бизнес-допущениям: стоимости контакта, "
        "вероятности успешного удержания и доле ценности клиента, которую можно сохранить."
    )
    doc.add_paragraph(
        "Осторожный сценарий использует более строгие допущения: ниже вероятность успеха и выше стоимость контакта. "
        "Базовый сценарий — центральная оценка. Оптимистичный сценарий показывает результат при более благоприятных "
        "допущениях. Параметр margin отвечает за сохраняемую долю ценности клиента, cost — за стоимость удержания, "
        "success — за вероятность успешного удержания."
    )
    scenario_df = pd.DataFrame(scenario_rows or [])
    scenario_cols = [
        c
        for c in ["scenario", "margin", "cost", "success", "clients_with_positive_ev", "best_k", "max_profit", "total_positive_ev"]
        if c in scenario_df.columns
    ]
    add_table(scenario_df.fillna("").to_dict(orient="records"), scenario_cols)
    scenario_chart = _save_scenario_max_effect_chart(scenario_df, Path(out_path).parent / "_docx_charts" / "scenario_max_profit.png")
    _add_picture_if_exists(doc, scenario_chart, width_inch=6.5)

    add_heading("3. Топ клиентов базового сценария", level=1)
    doc.add_paragraph("Первые строки приоритетного списка показывают клиентов, с которых разумно начинать кампанию удержания.")
    top_df = pd.DataFrame(top_clients_preview or [])
    top_cols = [
        c
        for c in ["entity_id", "risk_segment", "priority", "recommended_action", "p_calibrated", "value_proxy", "EV_base", "reason_1"]
        if c in top_df.columns
    ]
    add_table(top_df.fillna("").head(5).to_dict(orient="records"), top_cols)

    add_heading("4. Сопоставление колонок", level=1)
    doc.add_paragraph("Ниже указано, какие колонки нового CSV были сопоставлены с колонками обучающего набора.")
    add_table(
        [{"Колонка в обучении": k, "Колонка в прогнозном CSV": v} for k, v in score_mapping.items()],
        ["Колонка в обучении", "Колонка в прогнозном CSV"],
    )

    add_heading("5. Файлы в архиве", level=1)
    file_rows = [
        {
            "Файл": "scored_clients.csv",
            "Что внутри": "Полный результат скоринга по всем объектам: вероятности, сегменты риска, EV и пояснения.",
        },
        {
            "Файл": "scenario_summary.csv",
            "Что внутри": "Сводка по экономическим сценариям: best_k, max_profit и число клиентов с положительным EV.",
        },
        {
            "Файл": "retention_priority_list_conservative.csv",
            "Что внутри": "Приоритетный список для осторожного сценария.",
        },
        {
            "Файл": "retention_priority_list_base.csv",
            "Что внутри": "Приоритетный список для базового сценария.",
        },
        {
            "Файл": "retention_priority_list_optimistic.csv",
            "Что внутри": "Приоритетный список для оптимистичного сценария.",
        },
    ]
    add_table(file_rows, ["Файл", "Что внутри"])

    add_heading("6. Словарь колонок CSV", level=1)
    for filename, path in csv_paths.items():
        try:
            cols = list(pd.read_csv(path, nrows=0).columns)
        except Exception:
            cols = []
        if not cols:
            continue
        add_heading(filename, level=2)
        add_table(_csv_column_dictionary(cols), ["Колонка", "Описание"])

    add_heading("7. Как читать результат", level=1)
    doc.add_paragraph(
        "Чем выше p_calibrated, тем выше риск оттока. Чем выше EV_base или scenario_ev, тем выгоднее включить клиента "
        "в удерживающую кампанию. Дополнительные колонки, которые не были сопоставлены с обучающим набором, "
        "не участвуют в расчёте прогноза."
    )

    out_p = Path(out_path)
    out_p.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out_p)
    return {"docx_path": str(out_p)}
